import csv
import os
import uuid
from datetime import datetime
from pathlib import Path

import xlrd
from django.conf import settings
from django.core.files import File
from openpyxl import load_workbook

from atendimento.models import Atendimento
from atendimento.services import registrar_historico
from core.utils import compose_import_key, parse_time_string
from importacao.models import ImportacaoPlanilha, RegistroImportado


EXPECTED_HEADERS = {
    "hora": "hora",
    "nome": "nome",
    "eleitor": "nome",
    "status": "status",
    "guichê": "guiche",
    "guiche": "guiche",
    "observação": "observacao",
    "observacao": "observacao",
    "cpf/titulo": "documento",
    "cpf": "documento",
    "titulo": "documento",
    "email": "email",
    "telefone": "telefone",
    "tipo de serviço": "tipo_servico",
    "tipo de servico": "tipo_servico",
}

ADVANCED_STATUS = {Atendimento.Status.CHAMADO, Atendimento.Status.ATENDIDO}
TEMP_IMPORT_DIR = Path(settings.MEDIA_ROOT) / "tmp_imports"


def _normalize_text(value):
    return str(value or "").strip()


def _normalize_headers(headers):
    return [EXPECTED_HEADERS.get(_normalize_text(header).lower(), _normalize_text(header).lower()) for header in headers]


def _is_total_row(row):
    first = _normalize_text(row.get("hora") or row.get("nome")).lower()
    return first.startswith("total")


def _normalize_row(row):
    return {EXPECTED_HEADERS.get(_normalize_text(key).lower(), _normalize_text(key).lower()): value for key, value in row.items()}


def _parse_portal_date(value):
    raw = _normalize_text(value)
    if not raw:
        return None
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


def save_temp_upload(uploaded_file):
    TEMP_IMPORT_DIR.mkdir(parents=True, exist_ok=True)
    suffix = Path(uploaded_file.name).suffix.lower()
    temp_name = f"{uuid.uuid4().hex}{suffix}"
    temp_path = TEMP_IMPORT_DIR / temp_name
    with temp_path.open("wb") as destination:
        for chunk in uploaded_file.chunks():
            destination.write(chunk)
    return str(temp_path)


def _load_rows_and_metadata(file_path):
    suffix = Path(file_path).suffix.lower()
    metadata = {"sheet_name": None, "agenda_data": None, "headers": [], "row_count": 0}

    if suffix == ".csv":
        with open(file_path, "r", encoding="utf-8-sig", newline="") as handle:
            lines = handle.read().splitlines()
        reader = csv.DictReader(lines)
        rows = []
        metadata["headers"] = _normalize_headers(reader.fieldnames or [])
        for row in reader:
            normalized = _normalize_row(row)
            if normalized.get("nome") and not _is_total_row(normalized):
                rows.append(normalized)
        metadata["row_count"] = len(rows)
        return rows, metadata

    if suffix == ".xls":
        workbook = xlrd.open_workbook(file_path)
        sheet = workbook.sheet_by_index(0)
        metadata["sheet_name"] = sheet.name
        first_row = sheet.row_values(0) if sheet.nrows else []
        if len(first_row) > 1 and _normalize_text(first_row[0]).lower().startswith("agenda do dia"):
            metadata["agenda_data"] = _parse_portal_date(first_row[1])
        header_row = None
        for idx in range(sheet.nrows):
            normalized = _normalize_headers(sheet.row_values(idx))
            if "hora" in normalized and "nome" in normalized:
                header_row = idx
                metadata["headers"] = normalized
                break
        if header_row is None:
            raise ValueError("Nao foi possivel localizar o cabecalho da planilha .xls.")
        rows = []
        for idx in range(header_row + 1, sheet.nrows):
            row = dict(zip(metadata["headers"], sheet.row_values(idx)))
            if _normalize_text(row.get("nome")) and not _is_total_row(row):
                rows.append(row)
        metadata["row_count"] = len(rows)
        return rows, metadata

    workbook = load_workbook(file_path, read_only=True, data_only=True)
    sheet = workbook.active
    metadata["sheet_name"] = sheet.title
    rows_raw = list(sheet.iter_rows(values_only=True))
    if rows_raw and len(rows_raw[0]) > 1 and _normalize_text(rows_raw[0][0]).lower().startswith("agenda do dia"):
        metadata["agenda_data"] = _parse_portal_date(rows_raw[0][1])
    header_row = None
    for idx, candidate in enumerate(rows_raw):
        normalized = _normalize_headers(candidate)
        if "hora" in normalized and "nome" in normalized:
            header_row = idx
            metadata["headers"] = normalized
            break
    if header_row is None:
        raise ValueError("Nao foi possivel localizar o cabecalho da planilha .xlsx.")
    rows = []
    for raw in rows_raw[header_row + 1 :]:
        row = dict(zip(metadata["headers"], raw))
        if _normalize_text(row.get("nome")) and not _is_total_row(row):
            rows.append(row)
    metadata["row_count"] = len(rows)
    return rows, metadata


def preview_import_file(file_path, preview_limit=8):
    rows, metadata = _load_rows_and_metadata(file_path)
    preview_rows = []
    for row in rows[:preview_limit]:
        preview_rows.append(
            {
                "hora": _normalize_text(row.get("hora")),
                "nome": _normalize_text(row.get("nome")),
                "documento": _normalize_text(row.get("documento")),
                "email": _normalize_text(row.get("email")),
                "telefone": _normalize_text(row.get("telefone")),
                "tipo_servico": _normalize_text(row.get("tipo_servico")),
            }
        )
    return {"rows": preview_rows, "metadata": metadata}


def create_import_from_temp(temp_file, tenant, user, data_referencia):
    importacao = ImportacaoPlanilha.objects.create(tenant=tenant, data_referencia=data_referencia, importado_por=user)
    with open(temp_file, "rb") as handle:
        importacao.arquivo.save(Path(temp_file).name, File(handle), save=True)
    processar_importacao(importacao)
    if os.path.exists(temp_file):
        os.remove(temp_file)
    return importacao


def reprocessar_importacao(origem, user):
    nova = ImportacaoPlanilha.objects.create(
        tenant=origem.tenant,
        data_referencia=origem.data_referencia,
        importado_por=user,
    )
    with origem.arquivo.open("rb") as handle:
        nome = Path(origem.arquivo.name).name
        nova.arquivo.save(nome, File(handle), save=True)
    processar_importacao(nova)
    return nova


def processar_importacao(importacao):
    vistos = set()
    tenant = importacao.tenant
    arquivo = importacao.arquivo
    arquivo.open("rb")
    temp_path = None
    try:
        temp_path = save_temp_upload(arquivo)
        rows, _metadata = _load_rows_and_metadata(temp_path)
        for index, row in enumerate(rows, start=1):
            nome = _normalize_text(row.get("nome"))
            if not nome:
                continue
            hora = parse_time_string(row.get("hora")) if row.get("hora") else None
            observacao = _normalize_text(row.get("observacao"))
            status_origem = _normalize_text(row.get("status"))
            guiche_origem = _normalize_text(row.get("guiche"))
            documento = _normalize_text(row.get("documento"))
            email = _normalize_text(row.get("email"))
            telefone = _normalize_text(row.get("telefone"))
            tipo_servico = _normalize_text(row.get("tipo_servico"))
            detail_note = tipo_servico or observacao
            import_key = compose_import_key(nome, hora, importacao.data_referencia)
            if import_key in vistos:
                RegistroImportado.objects.create(importacao=importacao, linha=index, nome=nome, hora=hora, status_origem=status_origem, guiche_origem=guiche_origem, observacao=detail_note, resultado=RegistroImportado.Resultado.IGNORADO, detalhe="Duplicidade no mesmo arquivo.")
                importacao.ignorados += 1
                continue
            vistos.add(import_key)
            atendimento = Atendimento.objects.filter(tenant=tenant, import_key=import_key).first()
            if atendimento:
                if atendimento.status in ADVANCED_STATUS:
                    RegistroImportado.objects.create(importacao=importacao, linha=index, nome=nome, hora=hora, status_origem=status_origem, guiche_origem=guiche_origem, observacao=detail_note, resultado=RegistroImportado.Resultado.PRESERVADO, detalhe="Status avancado preservado.")
                    importacao.preservados += 1
                    continue
                changed = False
                for field, value in {"observacao": observacao, "hora_agendada": hora, "documento": documento, "email": email, "telefone": telefone, "tipo_servico": tipo_servico}.items():
                    if value and getattr(atendimento, field) != value:
                        setattr(atendimento, field, value)
                        changed = True
                if changed:
                    atendimento.save()
                    registrar_historico(atendimento, importacao.importado_por, "reimportacao", "Registro atualizado pela planilha.")
                    resultado = RegistroImportado.Resultado.ATUALIZADO
                    importacao.atualizados += 1
                else:
                    resultado = RegistroImportado.Resultado.PRESERVADO
                    importacao.preservados += 1
                RegistroImportado.objects.create(importacao=importacao, linha=index, nome=nome, hora=hora, status_origem=status_origem, guiche_origem=guiche_origem, observacao=detail_note, resultado=resultado, detalhe="Registro existente tratado de forma incremental.")
                continue
            atendimento = Atendimento.objects.create(tenant=tenant, nome=nome, documento=documento, email=email, telefone=telefone, tipo_servico=tipo_servico, data_referencia=importacao.data_referencia, hora_agendada=hora, status=Atendimento.Status.AGUARDANDO, classificacao_fila=Atendimento.Classificacao.NO_HORARIO, observacao=observacao, origem_importacao=importacao)
            registrar_historico(atendimento, importacao.importado_por, "importacao", "Registro inserido pela importacao inicial.")
            RegistroImportado.objects.create(importacao=importacao, linha=index, nome=nome, hora=hora, status_origem=status_origem, guiche_origem=guiche_origem, observacao=detail_note, resultado=RegistroImportado.Resultado.INSERIDO, detalhe="Novo agendamento importado.")
            importacao.inseridos += 1
    finally:
        arquivo.close()
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)
    importacao.save(update_fields=["inseridos", "atualizados", "preservados", "ignorados"])
    return importacao
